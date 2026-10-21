"""Recalcula resultados/pesos.xlsx con LibreOffice y comprueba que las fórmulas de Excel
den las mismas probabilidades y clases que la red en Python.

    python verificar_excel.py
"""

import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ARCHIVO = Path(__file__).resolve().parent / "resultados" / "pesos.xlsx"


def main():
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(ARCHIVO)],
                       check=True, capture_output=True)
        ws = load_workbook(Path(tmp) / ARCHIVO.name, data_only=True)["Comprobacion"]
    encabezados = [c.value for c in ws[6]]
    col = {n: i for i, n in enumerate(encabezados)}
    filas = [r for r in ws.iter_rows(min_row=7, values_only=True) if r[0] is not None]
    iguales = sum(r[col["¿Excel = Python?"]] == "Sí" for r in filas)
    dif = max(r[col["Dif. máx. p"]] for r in filas)
    print(f"filas recalculadas: {len(filas)}")
    print(f"Excel = Python: {iguales} de {len(filas)}  (diferencia máxima en p: {dif:.2e})")
    print(f"exactitud entrenamiento en Excel: {ws['D3'].value:.2%}")
    print(f"exactitud prueba en Excel:        {ws['D4'].value:.2%}")
    sys.exit(0 if iguales == len(filas) else 1)


if __name__ == "__main__":
    main()
