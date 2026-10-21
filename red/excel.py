"""Exporta los pesos a Excel con una hoja que recalcula la red usando fórmulas.

La hoja "Comprobacion" no copia los resultados de Python: los recalcula a partir de
las celdas de la hoja "Pesos finales". Si alguien cambia un peso, cambian las
predicciones, lo que permite comprobar a mano que los pesos generan el resultado.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as col

from .datos import CARACTERISTICAS, CLASES

NEGRITA = Font(bold=True)
TITULO = Font(bold=True, size=14)
ENCABEZADO = PatternFill("solid", fgColor="DDE7F0")
FORMATO = "0.000000"
HOJA_PESOS = "Pesos finales"


def _ref(hoja, fila, columna):
    return f"'{hoja}'!${col(columna)}${fila}"


def _encabezado(ws, fila, columna, texto):
    celda = ws.cell(fila, columna, texto)
    celda.font = NEGRITA
    celda.fill = ENCABEZADO
    celda.alignment = Alignment(horizontal="center")


def _matriz(ws, fila, nombre, M, nombres_filas, nombres_columnas):
    """Escribe M con encabezados. Regresa refs[i][j] absolutas a cada celda y la fila siguiente libre."""
    ws.cell(fila, 1, nombre).font = NEGRITA
    for j, n in enumerate(nombres_columnas):
        _encabezado(ws, fila + 1, j + 2, n)
    refs = []
    for i, n in enumerate(nombres_filas):
        ws.cell(fila + 2 + i, 1, n).font = NEGRITA
        fila_refs = []
        for j in range(M.shape[1]):
            c = ws.cell(fila + 2 + i, j + 2, float(M[i, j]))
            c.number_format = FORMATO
            fila_refs.append(_ref(ws.title, fila + 2 + i, j + 2))
        refs.append(fila_refs)
    return refs, fila + 3 + len(nombres_filas)


def _hoja_pesos(wb, titulo, pesos, est=None):
    ws = wb.create_sheet(titulo)
    ws.cell(1, 1, titulo).font = TITULO
    n_ent, n_ocu = pesos["W1"].shape
    n_sal = pesos["W2"].shape[1]
    ocultas = [f"h{k + 1}" for k in range(n_ocu)]
    salidas = [f"o{m + 1} ({c})" for m, c in enumerate(CLASES)]
    refs = {}
    fila = 3
    refs["W1"], fila = _matriz(ws, fila, f"W1 ({n_ent} x {n_ocu}), entrada → oculta", pesos["W1"], list(CARACTERISTICAS), ocultas)
    refs["b1"], fila = _matriz(ws, fila, f"b1 ({n_ocu})", pesos["b1"][None, :], ["b1"], ocultas)
    refs["W2"], fila = _matriz(ws, fila, f"W2 ({n_ocu} x {n_sal}), oculta → salida", pesos["W2"], ocultas, salidas)
    refs["b2"], fila = _matriz(ws, fila, f"b2 ({n_sal})", pesos["b2"][None, :], ["b2"], salidas)
    if est is not None:
        refs["media"], fila = _matriz(ws, fila, "Estandarización z = (x - media) / desviación", est.media[None, :], ["media"], list(CARACTERISTICAS))
        fila_desv = fila - 1  # justo debajo de la fila de la media
        refs["desv"] = [[_ref(ws.title, fila_desv, j + 2) for j in range(n_ent)]]
        ws.cell(fila_desv, 1, "desviación").font = NEGRITA
        for j, v in enumerate(est.desv):
            ws.cell(fila_desv, j + 2, float(v)).number_format = FORMATO
    ws.column_dimensions["A"].width = 16
    for j in range(2, max(n_ocu, n_ent) + 2):
        ws.column_dimensions[col(j)].width = 19
    return refs


def _argmax(r, c_p, n_sal):
    """Fórmula con IF anidados que regresa el nombre de la clase con mayor probabilidad."""
    p = [f"{col(c_p + m)}{r}" for m in range(n_sal)]
    formula = f'"{CLASES[-1]}"'
    for m in range(n_sal - 2, -1, -1):
        mayor = ",".join(f"{p[m]}>={p[k]}" for k in range(m + 1, n_sal))
        formula = f'IF(AND({mayor}),"{CLASES[m]}",{formula})'
    return "=" + formula


def _hoja_comprobacion(wb, refs, X, y, conjunto, P_python, pred_python):
    ws = wb.create_sheet("Comprobacion")
    n_ent = X.shape[1]
    n_ocu = len(refs["b1"][0])
    n_sal = len(refs["b2"][0])

    columnas = (["#", "Conjunto"] + [f"x{j + 1} {c}" for j, c in enumerate(CARACTERISTICAS)]
                + [f"z{j + 1}" for j in range(n_ent)] + [f"h{k + 1}" for k in range(n_ocu)]
                + [f"o{m + 1}" for m in range(n_sal)] + [f"p{m + 1} Excel" for m in range(n_sal)]
                + ["Predicción Excel", "Clase real", "¿Correcta?"]
                + [f"p{m + 1} Python" for m in range(n_sal)]
                + ["Predicción Python", "Dif. máx. p", "¿Excel = Python?"])
    c_x, c_z = 3, 3 + n_ent
    c_h = c_z + n_ent
    c_o = c_h + n_ocu
    c_p = c_o + n_sal
    c_pred, c_real, c_ok = c_p + n_sal, c_p + n_sal + 1, c_p + n_sal + 2
    c_ppy = c_ok + 1
    c_predpy, c_dif, c_igual = c_ppy + n_sal, c_ppy + n_sal + 1, c_ppy + n_sal + 2

    f0 = 7  # primera fila de datos
    fN = f0 + len(X) - 1
    ws.cell(1, 1, "Comprobación de la red con fórmulas de Excel").font = TITULO
    ws.cell(2, 1, "Cada fila recalcula z, h = tanh(z·W1 + b1), o = h·W2 + b2 y p = softmax(o) con los pesos de la hoja 'Pesos finales'.")
    resumen = [
        ("Exactitud entrenamiento", f'=COUNTIFS(B{f0}:B{fN},"entrenamiento",{col(c_ok)}{f0}:{col(c_ok)}{fN},"Sí")/COUNTIF(B{f0}:B{fN},"entrenamiento")'),
        ("Exactitud prueba", f'=COUNTIFS(B{f0}:B{fN},"prueba",{col(c_ok)}{f0}:{col(c_ok)}{fN},"Sí")/COUNTIF(B{f0}:B{fN},"prueba")'),
        ("Filas donde Excel = Python", f'=COUNTIF({col(c_igual)}{f0}:{col(c_igual)}{fN},"Sí")&" de "&COUNTA(A{f0}:A{fN})'),
    ]
    for i, (texto, formula) in enumerate(resumen):
        ws.cell(3 + i, 1, texto).font = NEGRITA
        c = ws.cell(3 + i, 4, formula)
        c.font = NEGRITA
        if i < 2:
            c.number_format = "0.00%"

    for j, n in enumerate(columnas):
        _encabezado(ws, f0 - 1, j + 1, n)

    for i in range(len(X)):
        r = f0 + i
        ws.cell(r, 1, i + 1)
        ws.cell(r, 2, conjunto[i])
        for j in range(n_ent):
            ws.cell(r, c_x + j, float(X[i, j]))
            ws.cell(r, c_z + j, f"=({col(c_x + j)}{r}-{refs['media'][0][j]})/{refs['desv'][0][j]}")
        for k in range(n_ocu):
            suma = "+".join(f"{col(c_z + j)}{r}*{refs['W1'][j][k]}" for j in range(n_ent))
            ws.cell(r, c_h + k, f"=TANH({suma}+{refs['b1'][0][k]})")
        for m in range(n_sal):
            suma = "+".join(f"{col(c_h + k)}{r}*{refs['W2'][k][m]}" for k in range(n_ocu))
            ws.cell(r, c_o + m, f"={suma}+{refs['b2'][0][m]}")
        exps = "+".join(f"EXP({col(c_o + m)}{r})" for m in range(n_sal))
        for m in range(n_sal):
            ws.cell(r, c_p + m, f"=EXP({col(c_o + m)}{r})/({exps})")
        ws.cell(r, c_pred, _argmax(r, c_p, n_sal))
        ws.cell(r, c_real, CLASES[y[i]])
        ws.cell(r, c_ok, f'=IF({col(c_pred)}{r}={col(c_real)}{r},"Sí","No")')
        for m in range(n_sal):
            ws.cell(r, c_ppy + m, float(P_python[i, m]))
        ws.cell(r, c_predpy, CLASES[pred_python[i]])
        difs = ",".join(f"ABS({col(c_p + m)}{r}-{col(c_ppy + m)}{r})" for m in range(n_sal))
        ws.cell(r, c_dif, f"=MAX({difs})").number_format = "0.0E+00"
        ws.cell(r, c_igual, f'=IF(AND({col(c_pred)}{r}={col(c_predpy)}{r},{col(c_dif)}{r}<1E-9),"Sí","No")')
        for c in range(c_z, c_pred):
            ws.cell(r, c).number_format = FORMATO
        for c in range(c_ppy, c_predpy):
            ws.cell(r, c).number_format = FORMATO

    for j in range(1, len(columnas) + 1):
        ws.column_dimensions[col(j)].width = 11
    for c in (2, c_pred, c_real, c_predpy):
        ws.column_dimensions[col(c)].width = 16
    ws.freeze_panes = ws.cell(f0, 3)


def exportar(ruta, red_cfg, pesos_ini, pesos_fin, est, X, y, conjunto, P_python, pred_python, historial, metricas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.cell(1, 1, "Red neuronal desde cero en TensorFlow, clasificación de Iris").font = TITULO
    filas = [
        ("Autor", "Daniel Alfredo Barreras Meraz (A01254805)"),
        ("Curso", "TC3002B.570 Desarrollo de aplicaciones avanzadas de ciencias computacionales"),
        ("Arquitectura", f"{red_cfg['entradas']} entradas → {red_cfg['ocultas']} ocultas (tanh) → {red_cfg['salidas']} salidas (softmax)"),
        ("Pérdida", "Entropía cruzada"),
        ("Optimización", f"Descenso de gradiente con lote completo, tasa {red_cfg['tasa']}, {red_cfg['epocas']} épocas"),
        ("Inicialización", f"Glorot uniforme, semilla {red_cfg['semilla']} (b = 0)"),
        ("Datos", f"Iris (UCI), {len(X)} flores; {conjunto.count('entrenamiento')} entrenamiento, {conjunto.count('prueba')} prueba"),
        ("Exactitud entrenamiento (Python)", metricas["exactitud_entrenamiento"]),
        ("Exactitud prueba (Python)", metricas["exactitud_prueba"]),
        ("Pérdida final", metricas["perdida_final"]),
        ("", ""),
        ("Hojas", ""),
        ("Pesos finales", "W1, b1, W2, b2 después del entrenamiento, más la media y desviación usadas para estandarizar"),
        ("Pesos iniciales", "Los mismos parámetros antes del entrenamiento"),
        ("Comprobacion", "Recalcula las 150 predicciones con fórmulas que leen 'Pesos finales' y las compara con Python"),
        ("Historial", "Pérdida de entrenamiento por época"),
    ]
    for i, (a, b) in enumerate(filas):
        ws.cell(3 + i, 1, a).font = NEGRITA
        c = ws.cell(3 + i, 2, b)
        c.alignment = Alignment(horizontal="left")
        if isinstance(b, float):
            c.number_format = "0.00%" if "Exactitud" in a else FORMATO
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 95

    refs = _hoja_pesos(wb, HOJA_PESOS, pesos_fin, est)
    _hoja_pesos(wb, "Pesos iniciales", pesos_ini)
    _hoja_comprobacion(wb, refs, X, y, conjunto, P_python, pred_python)

    wh = wb.create_sheet("Historial")
    _encabezado(wh, 1, 1, "Época")
    _encabezado(wh, 1, 2, "Pérdida")
    for i, (e, l) in enumerate(historial):
        wh.cell(2 + i, 1, e)
        wh.cell(2 + i, 2, l).number_format = FORMATO
    wh.column_dimensions["B"].width = 14
    wb.save(ruta)
